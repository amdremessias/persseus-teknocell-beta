import re, json, csv, unicodedata, difflib, openpyxl
from collections import defaultdict, Counter
from openpyxl.styles import Font, PatternFill

def cnorm(s):
    s = unicodedata.normalize('NFKD', s).encode('ascii', 'ignore').decode()
    return re.sub(r'\s+', ' ', s.upper()).strip()

def phone_keys(tel):
    d = re.sub(r'\D', '', tel)
    if d.startswith('55'):
        d = d[2:]
    keys = {d}
    if len(d) == 11 and d[2] == '9':
        keys.add(d[:2] + d[3:])
    if len(d) == 10:
        keys.add(d[:2] + '9' + d[2:])
    return keys

# 1. cadastro de clientes (nome -> telefone)
clientes = json.load(open('docs/_clientes.json'))
cli_exact = {}                 # nome_norm -> telefone
cli_fl = defaultdict(set)      # (primeiro,ultimo) -> {telefone}
cli_all = []                   # (nome_norm, telefone)
for c in clientes:
    nn = cnorm(c['nome']); tel = c['telefone']
    cli_exact.setdefault(nn, tel)
    cli_all.append((nn, tel))
    t = nn.split()
    if len(t) >= 2:
        cli_fl[(t[0], t[-1])].add(tel)

# 2. índice de leads por telefone
lead_by_phone = {}
with open('docs/_lead_phones.csv') as f:
    for row in csv.DictReader(f):
        for k in phone_keys(row['telefone']):
            lead_by_phone.setdefault(k, row['lead_id'])

def find_phone(nome_norm):
    if nome_norm in cli_exact:
        return cli_exact[nome_norm], 'cadastro_exato'
    t = nome_norm.split()
    if len(t) >= 2:
        s = cli_fl.get((t[0], t[-1]), set())
        if len(s) == 1:
            return next(iter(s)), 'cadastro_primeiro_ultimo'
        if len(s) > 1:
            return None, None
    scored = sorted(((difflib.SequenceMatcher(None, nome_norm, nn).ratio(), tel) for nn, tel in cli_all), key=lambda x: -x[0])
    if scored and scored[0][0] >= 0.90 and (len(scored) < 2 or scored[0][0] - scored[1][0] >= 0.06):
        return scored[0][1], f'cadastro_fuzzy_{scored[0][0]:.2f}'
    return None, None

# 3. atualiza as linhas
rows = list(csv.DictReader(open('docs/vendas-retroativas-v2.csv')))
ganhou, sem = 0, 0
for r in rows:
    if r['secao'] != '2-SEM-MATCH':
        continue
    tel, tp = find_phone(cnorm(r['nome']))
    if tel:
        lead_id = ''
        for k in phone_keys(tel):
            if k in lead_by_phone:
                lead_id = lead_by_phone[k]; break
        r['telefone'] = tel
        r['lead_id'] = lead_id
        r['match'] = tp + ('+lead' if lead_id else '+novo_lead')
        r['secao'] = '1-MATCH'
        ganhou += 1
    else:
        sem += 1

# 4. reordena e regrava CSV + XLSX
order = {'1-MATCH': 0, '2-SEM-MATCH': 1, '3-EXCLUIDO': 2, '4-NAO-APPLE': 3}
rows.sort(key=lambda d: (order[d['secao']], d['data_venda']))
cols = ['secao','nome','telefone','lead_id','produto','modelo','valor','data_venda','vendedor','armazenamento','seminovo','match','motivo_exclusao']
with open('docs/vendas-retroativas-v2.csv', 'w', newline='') as f:
    w = csv.DictWriter(f, fieldnames=cols); w.writeheader()
    for d in rows: w.writerow(d)

wb = openpyxl.Workbook(); wb.remove(wb.active)
TABS = [
 ('1-MATCHES (prontos)','1-MATCH',"Vendas Apple com telefone (lead do CRM ou cadastro da loja). lead_id vazio = Etapa 2 cria o lead. Danis: marca 'seminovo' se for.",'C7F5C7'),
 ('2-SEM MATCH','2-SEM-MATCH',"Vendas Apple ainda sem telefone (nome não achado nem no CRM nem no cadastro). Danis preenche manualmente quem quiser.",'FFF3C4'),
 ('3-EXCLUIDOS','3-EXCLUIDO',"B2B (Invictos), balcao (Consumidor/Cliente Padrao) e duplicatas. Só conferência.",'F5D0D0'),
 ('4-NAO-APPLE','4-NAO-APPLE',"Vendas não-Apple. Fora do pós-venda; só referência.",'E0E0E0'),
]
for tabname, sec, desc, color in TABS:
    ws = wb.create_sheet(tabname[:31])
    ws['A1'] = desc; ws['A1'].font = Font(italic=True, size=9); ws.merge_cells('A1:M1')
    for j, c in enumerate(cols, 1):
        cell = ws.cell(3, j, c); cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='444444')
    r = 4
    for row in [x for x in rows if x['secao'] == sec]:
        for j, c in enumerate(cols, 1):
            ws.cell(r, j, row[c] if c != 'secao' else '')
        r += 1
    for col, w in zip('ABCDEFGHIJKLM', [1, 34, 15, 26, 9, 20, 10, 12, 16, 14, 10, 24, 20]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A4'
wb.save('docs/vendas-retroativas-v2.xlsx')

# 5. report
a1 = sum(1 for r in rows if r['secao'] == '1-MATCH')
com_lead = sum(1 for r in rows if r['secao'] == '1-MATCH' and r['lead_id'])
sem_lead = a1 - com_lead
print(f"clientes no cadastro: {len(clientes)}")
print(f"dos 104 sem-match: ganharam telefone = {ganhou} | sobraram sem = {sem}")
print(f"origem dos matches novos: {dict(Counter(r['match'] for r in rows if r['secao']=='1-MATCH' and 'cadastro' in r['match']))}")
print(f"NOVA aba 1 (prontos): {a1}  ->  com lead_id existente: {com_lead} | sem lead_id (criar na Etapa 2): {sem_lead}")
print(f"aba 2 (ainda sem telefone): {sum(1 for r in rows if r['secao']=='2-SEM-MATCH')}")
